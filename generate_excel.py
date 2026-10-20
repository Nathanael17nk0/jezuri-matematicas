import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_mcd_workbook():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    font_title = Font(name="Calibri", size=16, bold=True, color="1E293B")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="64748B")
    font_section = Font(name="Calibri", size=13, bold=True, color="0F172A")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="1E293B")
    font_normal = Font(name="Calibri", size=11, color="1E293B")
    font_result = Font(name="Calibri", size=14, bold=True, color="065F46")
    font_note = Font(name="Calibri", size=10, italic=True, color="475569")

    fill_primary = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid") # Deep Blue
    fill_accent = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")  # Teal
    fill_purple = PatternFill(start_color="6D28D9", end_color="6D28D9", fill_type="solid")  # Purple
    fill_input = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")   # Soft yellow for input
    fill_result = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")  # Soft emerald
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")   # Light slate
    fill_highlight = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid") # Soft indigo

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    thick_bottom = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='medium', color='1E40AF')
    )
    double_bottom = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='double', color='0F172A')
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ==========================================
    # SHEET 1: MCD 2 NÚMEROS
    # ==========================================
    ws2 = wb.create_sheet(title="MCD 2 Números")
    ws2.views.sheetView[0].showGridLines = True

    # Title
    ws2["B2"] = "TABLA DINÁMICA: MÁXIMO COMÚN DIVISOR (2 NÚMEROS)"
    ws2["B2"].font = font_title
    ws2["B3"] = "Introduce los números en las celdas amarillas. Las fórmulas calcularán el MCD y los pasos."
    ws2["B3"].font = font_subtitle

    # Inputs
    ws2["B5"] = "Número A"
    ws2["B5"].font = font_bold
    ws2["C5"] = 48
    ws2["C5"].font = font_bold
    ws2["C5"].fill = fill_input
    ws2["C5"].alignment = align_center
    ws2["C5"].border = thin_border

    ws2["B6"] = "Número B"
    ws2["B6"].font = font_bold
    ws2["C6"] = 72
    ws2["C6"].font = font_bold
    ws2["C6"].fill = fill_input
    ws2["C6"].alignment = align_center
    ws2["C6"].border = thin_border

    # Direct Formula Result
    ws2["B8"] = "MCD (Fórmula Excel):"
    ws2["B8"].font = font_section
    ws2["C8"] = "=GCD(C5,C6)"
    ws2["C8"].font = font_result
    ws2["C8"].fill = fill_result
    ws2["C8"].alignment = align_center
    ws2["C8"].border = double_bottom

    ws2["D8"] = "Fórmula: =M.C.D(C5; C6) o =GCD(C5, C6)"
    ws2["D8"].font = font_note

    # Table 1: Descomposición Simultánea
    ws2["B11"] = "1. Tabla de Descomposición Simultánea en Factores Primos"
    ws2["B11"].font = font_section

    headers_t1 = ["Paso", "Valor A", "Valor B", "Divisor Primo Común", "Explicación"]
    for col_idx, h in enumerate(headers_t1, start=2):
        cell = ws2.cell(row=13, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_primary
        cell.alignment = align_center
        cell.border = thin_border

    # Sample rows pre-computed with formulas for 48 and 72
    rows_2num = [
        (1, 48, 72, 2, "Ambos son divisibles entre 2"),
        (2, 24, 36, 2, "Ambos son divisibles entre 2"),
        (3, 12, 18, 2, "Ambos son divisibles entre 2"),
        (4, 6, 9, 3, "Ambos son divisibles entre 3"),
        ("Fin", 2, 3, "-", "Los cocientes 2 y 3 ya no tienen divisores comunes (coprimos)")
    ]

    for r_idx, r_data in enumerate(rows_2num, start=14):
        for c_idx, val in enumerate(r_data, start=2):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_bold if r_idx == 18 else font_normal
            cell.alignment = align_center if c_idx < 6 else align_left
            cell.border = thin_border
            if r_idx % 2 == 1:
                cell.fill = fill_zebra
            if c_idx == 5 and val != "-":
                cell.fill = fill_highlight

    ws2["B20"] = "Cálculo del MCD:"
    ws2["B20"].font = font_bold
    ws2["C20"] = "=PRODUCT(E14:E17)"
    ws2["C20"].font = font_result
    ws2["C20"].fill = fill_result
    ws2["C20"].alignment = align_center
    ws2["C20"].border = double_bottom
    ws2["D20"] = "Multiplicación de los factores primos comunes: 2 × 2 × 2 × 3 = 24"
    ws2["D20"].font = font_bold

    # Table 2: Algoritmo de Euclides
    ws2["B23"] = "2. Tabla del Algoritmo de Euclides (Divisiones Sucesivas)"
    ws2["B23"].font = font_section

    # Layout Euclides:
    # Row 25: Cocientes (q)
    # Row 26: Dividendos / Divisores
    # Row 27: Residuos (r)
    ws2["B25"] = "Cocientes (q)"
    ws2["B25"].font = font_bold
    ws2["B25"].fill = fill_zebra
    ws2["B25"].border = thin_border

    ws2["B26"] = "Números / Divisores"
    ws2["B26"].font = font_bold
    ws2["B26"].fill = fill_primary
    ws2["B26"].font = font_header
    ws2["B26"].border = thin_border

    ws2["B27"] = "Residuos (r)"
    ws2["B27"].font = font_bold
    ws2["B27"].fill = fill_zebra
    ws2["B27"].border = thin_border

    # Paso 1: 72 / 48 -> q=1, r=24
    # Paso 2: 48 / 24 -> q=2, r=0 (MCD = 24)
    ws2["C25"] = "-"
    ws2["D25"] = 1
    ws2["E25"] = 2
    ws2["C26"] = 72
    ws2["D26"] = 48
    ws2["E26"] = 24
    ws2["C27"] = 24
    ws2["D27"] = 0
    ws2["E27"] = "-"

    for c in ["C", "D", "E"]:
        ws2[f"{c}25"].alignment = align_center
        ws2[f"{c}25"].border = thin_border
        ws2[f"{c}26"].alignment = align_center
        ws2[f"{c}26"].font = font_bold
        ws2[f"{c}26"].border = thin_border
        ws2[f"{c}27"].alignment = align_center
        ws2[f"{c}27"].border = thin_border

    ws2["E26"].fill = fill_result
    ws2["B29"] = "Resultado por Euclides:"
    ws2["B29"].font = font_bold
    ws2["C29"] = 24
    ws2["C29"].font = font_result
    ws2["C29"].fill = fill_result
    ws2["C29"].alignment = align_center
    ws2["C29"].border = double_bottom
    ws2["D29"] = "El último divisor con residuo cero es el MCD = 24"
    ws2["D29"].font = font_note

    # ==========================================
    # SHEET 2: MCD 3 NÚMEROS
    # ==========================================
    ws3 = wb.create_sheet(title="MCD 3 Números")
    ws3.views.sheetView[0].showGridLines = True

    ws3["B2"] = "TABLA DINÁMICA: MÁXIMO COMÚN DIVISOR (3 NÚMEROS)"
    ws3["B2"].font = font_title
    ws3["B3"] = "Introduce 3 números para calcular la descomposición simultánea y el MCD compuesto."
    ws3["B3"].font = font_subtitle

    # Inputs
    inputs_3 = [("Número A", 60), ("Número B", 90), ("Número C", 150)]
    for idx, (label, val) in enumerate(inputs_3, start=5):
        ws3[f"B{idx}"] = label
        ws3[f"B{idx}"].font = font_bold
        ws3[f"C{idx}"] = val
        ws3[f"C{idx}"].font = font_bold
        ws3[f"C{idx}"].fill = fill_input
        ws3[f"C{idx}"].alignment = align_center
        ws3[f"C{idx}"].border = thin_border

    # MCD Direct Formula
    ws3["B9"] = "MCD (Fórmula Excel):"
    ws3["B9"].font = font_section
    ws3["C9"] = "=GCD(C5,C6,C7)"
    ws3["C9"].font = font_result
    ws3["C9"].fill = fill_result
    ws3["C9"].alignment = align_center
    ws3["C9"].border = double_bottom

    ws3["D9"] = "Fórmula: =M.C.D(C5; C6; C7) o =GCD(C5, C6, C7)"
    ws3["D9"].font = font_note

    # Simultaneous Decomposition Table
    ws3["B12"] = "1. Tabla de Descomposición Simultánea (3 Números)"
    ws3["B12"].font = font_section

    headers_t3 = ["Paso", "Valor A", "Valor B", "Valor C", "Divisor Primo Común", "Explicación"]
    for col_idx, h in enumerate(headers_t3, start=2):
        cell = ws3.cell(row=14, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_accent
        cell.alignment = align_center
        cell.border = thin_border

    rows_3num = [
        (1, 60, 90, 150, 2, "Los 3 números son pares, divisibles entre 2"),
        (2, 30, 45, 75, 3, "La suma de cifras de cada uno es múltiplo de 3 (divisibles entre 3)"),
        (3, 10, 15, 25, 5, "Todos terminan en 0 o 5 (divisibles entre 5)"),
        ("Fin", 2, 3, 5, "-", "Los cocientes 2, 3 y 5 ya no tienen divisores comunes simultáneos")
    ]

    for r_idx, r_data in enumerate(rows_3num, start=15):
        for c_idx, val in enumerate(r_data, start=2):
            cell = ws3.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_bold if r_idx == 18 else font_normal
            cell.alignment = align_center if c_idx < 7 else align_left
            cell.border = thin_border
            if r_idx % 2 == 1:
                cell.fill = fill_zebra
            if c_idx == 6 and val != "-":
                cell.fill = fill_highlight

    ws3["B20"] = "Cálculo del MCD:"
    ws3["B20"].font = font_bold
    ws3["C20"] = "=PRODUCT(F15:F17)"
    ws3["C20"].font = font_result
    ws3["C20"].fill = fill_result
    ws3["C20"].alignment = align_center
    ws3["C20"].border = double_bottom
    ws3["D20"] = "Factores comunes: 2 × 3 × 5 = 30"
    ws3["D20"].font = font_bold

    # Euclides encadenado para 3 números
    ws3["B23"] = "2. Algoritmo de Euclides Encadenado: MCD(A, B, C) = MCD(MCD(A, B), C)"
    ws3["B23"].font = font_section

    ws3["B25"] = "Paso 1: MCD(60, 90)"
    ws3["B25"].font = font_bold
    ws3["C25"] = "=GCD(C5,C6)"
    ws3["C25"].font = font_bold
    ws3["C25"].fill = fill_highlight
    ws3["C25"].alignment = align_center
    ws3["C25"].border = thin_border
    ws3["D25"] = "90 / 60 -> q=1, r=30; luego 60 / 30 -> q=2, r=0. Resultado intermedio = 30"
    ws3["D25"].font = font_note

    ws3["B27"] = "Paso 2: MCD(Resultado 1, C) = MCD(30, 150)"
    ws3["B27"].font = font_bold
    ws3["C27"] = "=GCD(C25,C7)"
    ws3["C27"].font = font_result
    ws3["C27"].fill = fill_result
    ws3["C27"].alignment = align_center
    ws3["C27"].border = double_bottom
    ws3["D27"] = "150 / 30 -> q=5, r=0. El MCD definitivo es 30"
    ws3["D27"].font = font_bold

    # ==========================================
    # SHEET 3: MCD 4 NÚMEROS
    # ==========================================
    ws4 = wb.create_sheet(title="MCD 4 Números")
    ws4.views.sheetView[0].showGridLines = True

    ws4["B2"] = "TABLA DINÁMICA: MÁXIMO COMÚN DIVISOR (4 NÚMEROS)"
    ws4["B2"].font = font_title
    ws4["B3"] = "Introduce 4 números para calcular la descomposición simultánea y el MCD compuesto."
    ws4["B3"].font = font_subtitle

    # Inputs
    inputs_4 = [("Número A", 120), ("Número B", 180), ("Número C", 240), ("Número D", 300)]
    for idx, (label, val) in enumerate(inputs_4, start=5):
        ws4[f"B{idx}"] = label
        ws4[f"B{idx}"].font = font_bold
        ws4[f"C{idx}"] = val
        ws4[f"C{idx}"].font = font_bold
        ws4[f"C{idx}"].fill = fill_input
        ws4[f"C{idx}"].alignment = align_center
        ws4[f"C{idx}"].border = thin_border

    # MCD Direct Formula
    ws4["B10"] = "MCD (Fórmula Excel):"
    ws4["B10"].font = font_section
    ws4["C10"] = "=GCD(C5,C6,C7,C8)"
    ws4["C10"].font = font_result
    ws4["C10"].fill = fill_result
    ws4["C10"].alignment = align_center
    ws4["C10"].border = double_bottom

    ws4["D10"] = "Fórmula: =M.C.D(C5; C6; C7; C8) o =GCD(C5, C6, C7, C8)"
    ws4["D10"].font = font_note

    # Simultaneous Decomposition Table
    ws4["B13"] = "1. Tabla de Descomposición Simultánea (4 Números)"
    ws4["B13"].font = font_section

    headers_t4 = ["Paso", "Valor A", "Valor B", "Valor C", "Valor D", "Divisor Primo Común", "Explicación"]
    for col_idx, h in enumerate(headers_t4, start=2):
        cell = ws4.cell(row=15, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_purple
        cell.alignment = align_center
        cell.border = thin_border

    rows_4num = [
        (1, 120, 180, 240, 300, 2, "Los 4 números son pares (divisibles entre 2)"),
        (2, 60, 90, 120, 150, 2, "Los 4 números son pares (divisibles entre 2)"),
        (3, 30, 45, 60, 75, 3, "Todos son múltiplos de 3"),
        (4, 10, 15, 20, 25, 5, "Todos terminan en 0 o 5 (divisibles entre 5)"),
        ("Fin", 2, 3, 4, 5, "-", "Los cocientes 2, 3, 4 y 5 ya no comparten divisor común a los cuatro")
    ]

    for r_idx, r_data in enumerate(rows_4num, start=16):
        for c_idx, val in enumerate(r_data, start=2):
            cell = ws4.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_bold if r_idx == 20 else font_normal
            cell.alignment = align_center if c_idx < 8 else align_left
            cell.border = thin_border
            if r_idx % 2 == 1:
                cell.fill = fill_zebra
            if c_idx == 7 and val != "-":
                cell.fill = fill_highlight

    ws4["B22"] = "Cálculo del MCD:"
    ws4["B22"].font = font_bold
    ws4["C22"] = "=PRODUCT(G16:G19)"
    ws4["C22"].font = font_result
    ws4["C22"].fill = fill_result
    ws4["C22"].alignment = align_center
    ws4["C22"].border = double_bottom
    ws4["D22"] = "Factores comunes: 2 × 2 × 3 × 5 = 60"
    ws4["D22"].font = font_bold

    # Euclides en cascada para 4 números
    ws4["B25"] = "2. Algoritmo de Euclides por Pares: MCD( MCD(A, B), MCD(C, D) )"
    ws4["B25"].font = font_section

    ws4["B27"] = "MCD Grupo 1: MCD(120, 180)"
    ws4["B27"].font = font_bold
    ws4["C27"] = "=GCD(C5,C6)"
    ws4["C27"].font = font_bold
    ws4["C27"].fill = fill_highlight
    ws4["C27"].alignment = align_center
    ws4["C27"].border = thin_border
    ws4["D27"] = "= 60"
    ws4["D27"].font = font_note

    ws4["B28"] = "MCD Grupo 2: MCD(240, 300)"
    ws4["B28"].font = font_bold
    ws4["C28"] = "=GCD(C7,C8)"
    ws4["C28"].font = font_bold
    ws4["C28"].fill = fill_highlight
    ws4["C28"].alignment = align_center
    ws4["C28"].border = thin_border
    ws4["D28"] = "= 60"
    ws4["D28"].font = font_note

    ws4["B30"] = "MCD Final: MCD(Grupo 1, Grupo 2)"
    ws4["B30"].font = font_bold
    ws4["C30"] = "=GCD(C27,C28)"
    ws4["C30"].font = font_result
    ws4["C30"].fill = fill_result
    ws4["C30"].alignment = align_center
    ws4["C30"].border = double_bottom
    ws4["D30"] = "= 60"
    ws4["D30"].font = font_bold

    # ==========================================
    # SHEET 4: MATRIZ DINÁMICA DE MCD
    # ==========================================
    ws_mat = wb.create_sheet(title="Matriz Dinámica Multi-Número")
    ws_mat.views.sheetView[0].showGridLines = True

    ws_mat["B2"] = "MATRIZ DINÁMICA DE MÁXIMO COMÚN DIVISOR"
    ws_mat["B2"].font = font_title
    ws_mat["B3"] = "Cruce de números para calcular el MCD entre cualquier par de valores automáticamente."
    ws_mat["B3"].font = font_subtitle

    matrix_numbers = [12, 18, 24, 30, 36, 48, 60, 72, 90, 120]

    # Matrix Headers
    ws_mat["B5"] = "MCD"
    ws_mat["B5"].font = font_header
    ws_mat["B5"].fill = fill_primary
    ws_mat["B5"].alignment = align_center
    ws_mat["B5"].border = thin_border

    for col_idx, num in enumerate(matrix_numbers, start=3):
        col_letter = get_column_letter(col_idx)
        cell = ws_mat[f"{col_letter}5"]
        cell.value = num
        cell.font = font_header
        cell.fill = fill_primary
        cell.alignment = align_center
        cell.border = thin_border

    for row_idx, num_r in enumerate(matrix_numbers, start=6):
        cell_r = ws_mat[f"B{row_idx}"]
        cell_r.value = num_r
        cell_r.font = font_header
        cell_r.fill = fill_primary
        cell_r.alignment = align_center
        cell_r.border = thin_border

        for col_idx, _ in enumerate(matrix_numbers, start=3):
            col_letter = get_column_letter(col_idx)
            cell_data = ws_mat[f"{col_letter}{row_idx}"]
            cell_data.value = f"=GCD($B{row_idx},{col_letter}$5)"
            cell_data.font = font_normal
            cell_data.alignment = align_center
            cell_data.border = thin_border
            if row_idx % 2 == 0:
                cell_data.fill = fill_zebra

    # ==========================================
    # SHEET 5: FACTORES PRIMOS PASO A PASO
    # ==========================================
    ws_fp = wb.create_sheet(title="Factores Primos Paso a Paso")
    ws_fp.views.sheetView[0].showGridLines = True

    ws_fp["B2"] = "DESCOMPOSICIÓN EN FACTORES PRIMOS (MÉTODO CANÓNICO)"
    ws_fp["B2"].font = font_title
    ws_fp["B3"] = "Tablas verticales de descomposición para cada número y selección del menor exponente común."
    ws_fp["B3"].font = font_subtitle

    # Cuatro tablas verticales lado a lado para 120, 180, 240, 300
    vert_examples = [
        ("Número A = 120", 2, [120, 60, 30, 15, 5, 1], [2, 2, 2, 3, 5, "-"], "120 = 2³ × 3¹ × 5¹"),
        ("Número B = 180", 6, [180, 90, 45, 15, 5, 1], [2, 2, 3, 3, 5, "-"], "180 = 2² × 3² × 5¹"),
        ("Número C = 240", 10, [240, 120, 60, 30, 15, 5, 1], [2, 2, 2, 2, 3, 5, "-"], "240 = 2⁴ × 3¹ × 5¹"),
        ("Número D = 300", 14, [300, 150, 75, 25, 5, 1], [2, 2, 3, 5, 5, "-"], "300 = 2² × 3¹ × 5²")
    ]

    for title_txt, start_c, quotients, primes, canon_str in vert_examples:
        c1_letter = get_column_letter(start_c)
        c2_letter = get_column_letter(start_c + 1)

        ws_fp[f"{c1_letter}5"] = title_txt
        ws_fp[f"{c1_letter}5"].font = font_bold
        ws_fp[f"{c1_letter}5"].fill = fill_highlight

        ws_fp[f"{c1_letter}6"] = "Cociente"
        ws_fp[f"{c1_letter}6"].font = font_header
        ws_fp[f"{c1_letter}6"].fill = fill_primary
        ws_fp[f"{c1_letter}6"].alignment = align_center

        ws_fp[f"{c2_letter}6"] = "Factor Primo"
        ws_fp[f"{c2_letter}6"].font = font_header
        ws_fp[f"{c2_letter}6"].fill = fill_accent
        ws_fp[f"{c2_letter}6"].alignment = align_center

        for idx, (q, p) in enumerate(zip(quotients, primes), start=7):
            cell_q = ws_fp[f"{c1_letter}{idx}"]
            cell_q.value = q
            cell_q.font = font_bold if q == 1 else font_normal
            cell_q.alignment = align_center
            cell_q.border = thin_border

            cell_p = ws_fp[f"{c2_letter}{idx}"]
            cell_p.value = p
            cell_p.font = font_bold if p != "-" else font_normal
            cell_p.alignment = align_center
            cell_p.border = thin_border

        last_r = 7 + len(quotients)
        ws_fp[f"{c1_letter}{last_r}"] = canon_str
        ws_fp[f"{c1_letter}{last_r}"].font = font_bold
        ws_fp[f"{c1_letter}{last_r}"].fill = fill_zebra

    # Matriz de menor exponente
    ws_fp["B17"] = "Matriz de Comparación: Factores Comunes con MENOR Exponente"
    ws_fp["B17"].font = font_section

    headers_comp = ["Número", "Forma Canónica", "Exp. de 2", "Exp. de 3", "Exp. de 5", "¿Factores Comunes?"]
    for idx_col, h_text in enumerate(headers_comp, start=2):
        cell_h = ws_fp.cell(row=19, column=idx_col, value=h_text)
        cell_h.font = font_header
        cell_h.fill = fill_purple
        cell_h.alignment = align_center
        cell_h.border = thin_border

    data_comp = [
        ("A (120)", "2³ × 3¹ × 5¹", 3, 1, 1, "Contiene 2, 3 y 5"),
        ("B (180)", "2² × 3² × 5¹", 2, 2, 1, "Contiene 2, 3 y 5"),
        ("C (240)", "2⁴ × 3¹ × 5¹", 4, 1, 1, "Contiene 2, 3 y 5"),
        ("D (300)", "2² × 3¹ × 5²", 2, 1, 2, "Contiene 2, 3 y 5"),
        ("M.C.D.", "2² × 3¹ × 5¹ = 60", "=MIN(D20:D23)", "=MIN(E20:E23)", "=MIN(F20:F23)", "MCD = 2² × 3¹ × 5¹ = 60")
    ]

    for r_idx, r_row in enumerate(data_comp, start=20):
        for c_idx, val in enumerate(r_row, start=2):
            cell = ws_fp.cell(row=r_idx, column=c_idx, value=val)
            is_mcd_row = r_idx == 24
            cell.font = font_bold if is_mcd_row else font_normal
            cell.alignment = align_center if c_idx > 2 else align_left
            cell.border = double_bottom if is_mcd_row else thin_border
            if is_mcd_row:
                cell.fill = fill_result

    # Pares de Factores Multiplicativos
    ws_fp["B27"] = "Tabla de Pares de Factores Multiplicativos (Ejemplo para Número A = 120)"
    ws_fp["B27"].font = font_section

    headers_pairs = ["Par #", "Factor A", "×", "Factor B", "=", "Producto", "Área Rectángulo"]
    for idx_c, hp in enumerate(headers_pairs, start=2):
        cell_p = ws_fp.cell(row=29, column=idx_c, value=hp)
        cell_p.font = font_header
        cell_p.fill = fill_accent
        cell_p.alignment = align_center
        cell_p.border = thin_border

    pairs_data = [
        (1, 1, "×", 120, "=", 120, "1 × 120"),
        (2, 2, "×", 60, "=", 120, "2 × 60"),
        (3, 3, "×", 40, "=", 120, "3 × 40"),
        (4, 4, "×", 30, "=", 120, "4 × 30"),
        (5, 5, "×", 24, "=", 120, "5 × 24"),
        (6, 6, "×", 20, "=", 120, "6 × 20"),
        (7, 8, "×", 15, "=", 120, "8 × 15"),
        (8, 10, "×", 12, "=", 120, "10 × 12")
    ]

    for r_idx, pr in enumerate(pairs_data, start=30):
        for c_idx, val in enumerate(pr, start=2):
            cell = ws_fp.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_normal
            cell.alignment = align_center
            cell.border = thin_border
            if r_idx % 2 == 1:
                cell.fill = fill_zebra

    # ==========================================
    # SHEET 6: GUÍA TEÓRICA Y FÓRMULAS
    # ==========================================
    ws_guide = wb.create_sheet(title="Guía y Fórmulas")
    ws_guide.views.sheetView[0].showGridLines = True

    ws_guide["B2"] = "GUÍA DE FÓRMULAS Y MÉTODOS DEL MÁXIMO COMÚN DIVISOR"
    ws_guide["B2"].font = font_title

    guide_content = [
        ("1. ¿Qué es el MCD?", "Es el mayor número entero que divide exactamente a dos o más números sin dejar residuo."),
        ("2. Fórmulas Nativas en Excel:", "En español: =M.C.D(núm1, núm2, [núm3], [núm4]...)\nEn inglés: =GCD(num1, num2, [num3], [num4]...)"),
        ("3. Método 1: Descomposición Simultánea:", "Se colocan los números juntos y se dividen ÚNICAMENTE por factores primos comunes a TODOS hasta que queden coprimos."),
        ("4. Método 2: Algoritmo de Euclides:", "Divisiones sucesivas: MCD(a, b) = MCD(b, a mod b) hasta residuo 0. Para 3 o 4 números: MCD(a,b,c) = MCD(MCD(a,b), c)."),
        ("5. Propiedad Fundamental:", "Si el MCD(a, b) = 1, los números se denominan 'primos entre sí' o 'coprimos'."),
        ("6. Relación con el MCM (para 2 números):", "a × b = MCD(a, b) × MCM(a, b)  ==>  MCM(a, b) = (a × b) / MCD(a, b)")
    ]

    for idx, (title, desc) in enumerate(guide_content, start=4):
        ws_guide[f"B{idx*2}"] = title
        ws_guide[f"B{idx*2}"].font = font_bold
        ws_guide[f"B{idx*2+1}"] = desc
        ws_guide[f"B{idx*2+1}"].font = font_normal

    # Column Width Auto-adjust
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if '\n' in val_str:
                    val_str = max(val_str.split('\n'), key=len)
                if len(val_str) > max_len and cell.row > 2:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        ws.column_dimensions["A"].width = 4
        ws.column_dimensions["B"].width = 24

    file_path = "/Users/air/Desktop/Proyectos/Jezuri/Tablas_Dinamicas_MCD.xlsx"
    wb.save(file_path)
    print(f"Archivo Excel generado con éxito en: {file_path}")

if __name__ == "__main__":
    create_mcd_workbook()
